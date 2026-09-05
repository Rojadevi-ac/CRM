import io
from datetime import datetime
from flask import Blueprint, request, jsonify, send_file
from services.report_service import ReportService
from middleware.auth_middleware import token_required
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

report_bp = Blueprint('reports', __name__, url_prefix='/api/reports')

@report_bp.route('/dashboard', methods=['GET'])
@token_required
def get_dashboard():
    kpis = ReportService.get_dashboard_kpis()
    return jsonify({'success': True, 'data': kpis}), 200

@report_bp.route('/leads', methods=['GET'])
@token_required
def get_lead_report():
    date_range = request.args.get('date_range', 'this_month')
    custom_from = request.args.get('custom_from')
    custom_to = request.args.get('custom_to')
    status = request.args.get('status')
    user_id = request.args.get('user_id')

    report = ReportService.get_lead_report(date_range, custom_from, custom_to, status, user_id)
    return jsonify({'success': True, 'data': report}), 200

@report_bp.route('/sales', methods=['GET'])
@token_required
def get_sales_report():
    date_range = request.args.get('date_range', 'this_month')
    custom_from = request.args.get('custom_from')
    custom_to = request.args.get('custom_to')
    stage = request.args.get('stage')
    owner_id = request.args.get('owner_id')

    report = ReportService.get_sales_report(date_range, custom_from, custom_to, stage, owner_id)
    return jsonify({'success': True, 'data': report}), 200

@report_bp.route('/salesperson', methods=['GET'])
@token_required
def get_salesperson_report():
    date_range = request.args.get('date_range', 'this_month')
    custom_from = request.args.get('custom_from')
    custom_to = request.args.get('custom_to')

    report = ReportService.get_salesperson_report(date_range, custom_from, custom_to)
    return jsonify({'success': True, 'data': report}), 200

@report_bp.route('/customers', methods=['GET'])
@token_required
def get_customer_report():
    date_range = request.args.get('date_range', 'this_month')
    custom_from = request.args.get('custom_from')
    custom_to = request.args.get('custom_to')

    report = ReportService.get_customer_report(date_range, custom_from, custom_to)
    return jsonify({'success': True, 'data': report}), 200

@report_bp.route('/export/excel', methods=['GET'])
@token_required
def export_excel():
    report_type = request.args.get('type', 'sales')
    date_range = request.args.get('date_range', 'this_month')
    custom_from = request.args.get('custom_from')
    custom_to = request.args.get('custom_to')

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    title_font = Font(name='Segoe UI', size=14, bold=True, color='1E3A8A')
    border_side = Side(border_style='thin', color='CBD5E1')
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    today_str = datetime.now().strftime('%Y-%m-%d')

    if report_type == 'sales' or report_type == 'all':
        ws = wb.create_sheet(title='Sales Summary')
        sales_data = ReportService.get_sales_report(date_range, custom_from, custom_to)
        
        ws['A1'] = "RD-CRM Sales Performance Report"
        ws['A1'].font = title_font
        ws['A2'] = f"Generated Date: {today_str} | Period: {sales_data['date_bounds']['from']} to {sales_data['date_bounds']['to']}"
        ws['A2'].font = Font(name='Segoe UI', size=9, italic=True)

        headers = ['Deal ID', 'Deal Name', 'Client/Company', 'Amount (INR)', 'Stage', 'Probability (%)', 'Weighted Value', 'Owner', 'Expected Close']
        ws.append([])
        ws.append(headers)

        header_row_idx = 4
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for d in sales_data['details']:
            ws.append([
                d['id'], d['deal_name'], d['client_name'], float(d['amount']), d['stage'],
                f"{d['probability']}%", float(d['weighted_value']), d['owner_name'], str(d['expected_closing_date'] or '')
            ])

    if report_type == 'leads' or report_type == 'all':
        ws_leads = wb.create_sheet(title='Leads Report')
        lead_data = ReportService.get_lead_report(date_range, custom_from, custom_to)

        ws_leads['A1'] = "RD-CRM Lead Analytics & Pipeline"
        ws_leads['A1'].font = title_font
        ws_leads['A2'] = f"Generated Date: {today_str} | Conversion Rate: {lead_data['conversion_rate']}%"
        ws_leads['A2'].font = Font(name='Segoe UI', size=9, italic=True)

        l_headers = ['Lead ID', 'Name', 'Email', 'Phone', 'Company', 'Source', 'Industry', 'Status', 'Priority', 'Est. Value (INR)', 'Assigned User', 'Created Date']
        ws_leads.append([])
        ws_leads.append(l_headers)

        header_row_idx = 4
        for col_idx, h in enumerate(l_headers, start=1):
            cell = ws_leads.cell(row=header_row_idx, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for l in lead_data['details']:
            ws_leads.append([
                l['id'], l['name'], l['email'], l['phone'], l['company_name'], l['source'],
                l['industry'], l['status'], l['priority'], float(l['estimated_value'] or 0),
                l['assigned_user'], str(l['created_date'])
            ])

    if report_type == 'salesperson' or report_type == 'all':
        ws_sp = wb.create_sheet(title='Salesperson Leaderboard')
        sp_data = ReportService.get_salesperson_report(date_range, custom_from, custom_to)

        ws_sp['A1'] = "RD-CRM Salesperson Performance Leaderboard"
        ws_sp['A1'].font = title_font
        ws_sp['A2'] = f"Generated Date: {today_str}"
        ws_sp['A2'].font = Font(name='Segoe UI', size=9, italic=True)

        sp_headers = ['Salesperson', 'Department', 'Assigned Leads', 'Qualified Leads', 'Total Deals', 'Won Deals', 'Lost Deals', 'Revenue (INR)', 'Conversion Rate (%)']
        ws_sp.append([])
        ws_sp.append(sp_headers)

        header_row_idx = 4
        for col_idx, h in enumerate(sp_headers, start=1):
            cell = ws_sp.cell(row=header_row_idx, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for p in sp_data['performance']:
            ws_sp.append([
                p['salesperson'], p['department'], p['assigned_leads'], p['qualified_leads'],
                p['total_deals'], p['won_deals'], p['lost_deals'], float(p['revenue']), f"{p['conversion_rate']}%"
            ])

    # Auto-adjust column widths for all sheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = str(cell.value or '')
                max_len = max(max_len, len(val))
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"RD_CRM_{report_type.capitalize()}_Report_{today_str}.xlsx"
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )
